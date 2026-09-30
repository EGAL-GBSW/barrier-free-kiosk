"""Run faster-whisper on the 41-file pilot ZIP and compare with the evaluation sheet."""

from __future__ import annotations

import argparse
import csv
import io
import json
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from zipfile import ZipFile, ZipInfo

from benchmark_stt import edit_distance, normalize, percentile
from stt_service import STTConfig, STTEngine


MAX_AUDIO_BYTES = 25 * 1024 * 1024
REQUIRED_COLUMNS = ("파일명", "원본 파일명", "정답 문장")
OUTPUT_COLUMNS = (
    "파일명", "녹음자", "문장번호", "STT ID", "정답 문장", "메뉴", "발화 유형",
    "원본 파일명", "ZIP 내부 경로", "STT 모델", "STT 추출 결과", "정확히 일치",
    "정규화 CER", "처리 시간(초)", "음성 길이(초)", "상태", "오류 메모",
)


@dataclass(frozen=True)
class PilotRow:
    values: dict[str, object]
    excel_row: int

    @property
    def original_name(self) -> str:
        return str(self.values["원본 파일명"]).strip()

    @property
    def reference(self) -> str:
        return str(self.values["정답 문장"]).strip()


def filename_key(name: str) -> str:
    """Make composed/decomposed Korean ZIP and Excel names comparable."""
    return unicodedata.normalize("NFC", PurePosixPath(name.replace("\\", "/")).name).casefold()


def read_pilot_rows(path: Path, sheet_name: str = "Pilot_Test") -> list[PilotRow]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install evaluation dependencies: python -m pip install -r requirements-dev.txt") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}")
        rows = workbook[sheet_name].iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise ValueError("Evaluation sheet is empty")
        columns = {str(value).strip(): index for index, value in enumerate(headers) if value is not None}
        missing = set(REQUIRED_COLUMNS) - columns.keys()
        if missing:
            raise ValueError(f"Missing evaluation columns: {', '.join(sorted(missing))}")

        records: list[PilotRow] = []
        seen: set[str] = set()
        for excel_row, cells in enumerate(rows, start=2):
            if not any(value is not None for value in cells):
                continue
            values = {
                column: cells[index] if index < len(cells) else None
                for column, index in columns.items()
            }
            record = PilotRow(values, excel_row)
            if not record.original_name or not normalize(record.reference):
                raise ValueError(f"Missing original filename/reference on Excel row {excel_row}")
            key = filename_key(record.original_name)
            if key in seen:
                raise ValueError(f"Duplicate original filename on Excel row {excel_row}: {record.original_name}")
            seen.add(key)
            records.append(record)
        if not records:
            raise ValueError("Evaluation sheet has no audio rows")
        return records
    finally:
        workbook.close()


def index_zip_audio(archive: ZipFile) -> dict[str, ZipInfo]:
    members: dict[str, ZipInfo] = {}
    for info in archive.infolist():
        if info.is_dir() or PurePosixPath(info.filename).suffix.lower() not in {".wav", ".mp3", ".m4a", ".webm"}:
            continue
        key = filename_key(info.filename)
        if key in members:
            raise ValueError(f"Duplicate audio filename in ZIP: {info.filename}")
        members[key] = info
    return members


def excel_trim(text: str) -> str:
    return " ".join(text.split())


def safe_csv_value(value: object) -> object:
    """Prevent source text or transcripts being interpreted as spreadsheet formulas."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def run(
    evaluation_xlsx: Path,
    audio_zip: Path,
    output_csv: Path,
    config: STTConfig,
    *,
    sheet_name: str = "Pilot_Test",
    limit: int | None = None,
    overwrite: bool = False,
) -> dict[str, object]:
    if limit is not None and limit < 1:
        raise ValueError("limit must be positive")
    summary_path = output_csv.with_suffix(".summary.json")
    if not overwrite and (output_csv.exists() or summary_path.exists()):
        raise FileExistsError(f"Output already exists; choose another path or use --overwrite: {output_csv}")

    records = read_pilot_rows(evaluation_xlsx, sheet_name)
    selected = records[:limit]
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    exact = errors = characters = succeeded = failed = missing_audio = 0
    times: list[float] = []
    load_seconds = 0.0

    with ZipFile(audio_zip) as archive:
        members = index_zip_audio(archive)
        expected_keys = {filename_key(record.original_name) for record in records}
        extra_audio = sorted(info.filename for key, info in members.items() if key not in expected_keys)
        engine = STTEngine(config)
        if any(filename_key(record.original_name) in members for record in selected):
            load_seconds = engine.load()

        with output_csv.open("w", newline="", encoding="utf-8-sig") as destination:
            writer = csv.DictWriter(destination, fieldnames=OUTPUT_COLUMNS)
            writer.writeheader()
            for index, record in enumerate(selected, start=1):
                row = {column: record.values.get(column) or "" for column in OUTPUT_COLUMNS}
                row["STT 모델"] = config.model_name
                member = members.get(filename_key(record.original_name))
                if member is None:
                    row["상태"] = "missing_audio"
                    row["오류 메모"] = "ZIP에서 원본 파일명을 찾지 못함"
                    missing_audio += 1
                else:
                    row["ZIP 내부 경로"] = member.filename
                    try:
                        if member.file_size > MAX_AUDIO_BYTES:
                            raise ValueError(f"Audio exceeds {MAX_AUDIO_BYTES} bytes")
                        result = engine.transcribe(io.BytesIO(archive.read(member)))
                        row["STT 추출 결과"] = result.text
                        row["상태"] = "ok"
                        row["처리 시간(초)"] = round(result.processing_seconds, 3)
                        row["음성 길이(초)"] = (
                            round(result.audio_duration_seconds, 3)
                            if result.audio_duration_seconds is not None else ""
                        )
                        row["정확히 일치"] = "O" if excel_trim(record.reference) == excel_trim(result.text) else "X"
                        expected = normalize(record.reference)
                        sample_errors = edit_distance(expected, normalize(result.text))
                        row["정규화 CER"] = round(sample_errors / len(expected), 4)
                        errors += sample_errors
                        characters += len(expected)
                        exact += row["정확히 일치"] == "O"
                        succeeded += 1
                        times.append(result.processing_seconds)
                    except Exception as exc:
                        row["상태"] = "error"
                        row["오류 메모"] = f"{type(exc).__name__}: {exc}"
                        failed += 1
                writer.writerow({key: safe_csv_value(value) for key, value in row.items()})
                print(
                    f"[{index}/{len(selected)}] {record.values.get('파일명', record.original_name)} "
                    f"| 정답: {record.reference} | 추출: {row['STT 추출 결과']} | {row['상태']}",
                    file=sys.stderr,
                )

    summary = {
        "model": config.model_name,
        "device": config.device,
        "compute_type": config.compute_type,
        "evaluation_rows": len(records),
        "processed": len(selected),
        "succeeded": succeeded,
        "failed": failed,
        "missing_audio": missing_audio,
        "extra_audio_in_zip": extra_audio,
        "exact_match_rate": round(exact / succeeded, 4) if succeeded else None,
        "normalized_cer": round(errors / characters, 4) if characters else None,
        "model_load_seconds": round(load_seconds, 3),
        "latency_p50_seconds": round(percentile(times, 0.5), 3) if times else None,
        "latency_p95_seconds": round(percentile(times, 0.95), 3) if times else None,
        "results_csv": str(output_csv),
    }
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--evaluation-xlsx", type=Path, required=True)
    parser.add_argument("--audio-zip", type=Path, required=True)
    parser.add_argument("--output-csv", type=Path, default=Path(__file__).parent / "benchmark_results" / "pilot_stt_results.csv")
    parser.add_argument("--sheet", default="Pilot_Test")
    parser.add_argument("--limit", type=int, help="Run only the first N evaluation rows")
    parser.add_argument("--overwrite", action="store_true")
    defaults = STTConfig.from_env()
    parser.add_argument("--model", default=defaults.model_name)
    parser.add_argument("--device", default=defaults.device)
    parser.add_argument("--compute-type", default=defaults.compute_type)
    args = parser.parse_args()
    summary = run(
        args.evaluation_xlsx, args.audio_zip, args.output_csv,
        STTConfig(args.model, args.device, args.compute_type),
        sheet_name=args.sheet, limit=args.limit, overwrite=args.overwrite,
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    if summary["failed"] or summary["missing_audio"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
